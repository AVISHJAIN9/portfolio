#!/usr/bin/env python3
import openpyxl
import os

for slot, fname, total_p0 in [
    ("04", "04_Columnar_Query_Engine_ULTIMATE_145_Master_STATUS.xlsx", 45),
    ("05", "05_Mini_QUIC_ULTIMATE_145_Master_STATUS.xlsx", 79),
    ("06", "06_LSM_KV_ENGINE_ULTIMATE_145_Master_STATUS.xlsx", 117)
]:
    p = os.path.join("evidence/workbooks", fname)
    if not os.path.exists(p):
        print(f"Workbook missing: {p}")
        continue
    wb = openpyxl.load_workbook(p, data_only=True)
    ws = wb["Implementation Status"]
    print(f"\n=======================================================")
    print(f"P0 TICK / CROSS STATUS SUMMARY — PROJECT {slot} ({fname})")
    print(f"=======================================================")
    header = f"{'ID':<8} | {'Imp':<3} | {'Ver':<3} | {'Status':<12} | {'Feature Name':<45} | Evidence"
    print(header)
    print("-" * 110)
    ver_count = 0
    p0_count = 0
    for r in range(12, ws.max_row + 1):
        fid = ws.cell(r, 1).value
        name = ws.cell(r, 2).value
        prio = ws.cell(r, 3).value
        imp = ws.cell(r, 4).value
        ver = ws.cell(r, 5).value
        stat = ws.cell(r, 6).value
        ev = ws.cell(r, 7).value or ""
        if fid and prio == "P0":
            p0_count += 1
            if ver == "✓":
                ver_count += 1
            n_str = str(name)[:45]
            print(f"{fid:<8} | {imp:<3} | {ver:<3} | {stat:<12} | {n_str:<45} | {ev}")
    print("-" * 110)
    print(f"P0 Summary for Project {slot}: Verified {ver_count}/{p0_count} (Expected Total P0: {total_p0})\n")
