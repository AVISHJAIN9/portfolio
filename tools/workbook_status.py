#!/usr/bin/env python3
"""
tools/workbook_status.py
Generates <workbook>_STATUS.xlsx under evidence/workbooks/ from master workbooks and features.jsonl.
Preserves original sheets, adds Feature ID / Implemented / Verified / Status / Evidence columns,
and appends a live formula-driven 'Implementation Status' rollup sheet.
"""

import os
import sys
import glob
import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from pathlib import Path
import subprocess
import datetime

ROOT_DIR = Path(__file__).resolve().parent.parent

GREEN_FILL = PatternFill(start_color="E2F0D9", end_color="E2F0D9", fill_type="solid")
GREEN_FONT = Font(name="Arial", size=10, bold=True, color="385723")
RED_FILL = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
RED_FONT = Font(name="Arial", size=10, bold=True, color="C65911")
HEADER_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
BORDER_THIN = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

def load_evidence_map():
    evidence_map = {}
    csv_path = ROOT_DIR / "evidence" / "feature_evidence.csv"
    if csv_path.exists():
        with open(csv_path) as f:
            for line in f.readlines()[1:]:
                parts = line.strip().split(",")
                if len(parts) >= 7:
                    feat_id = parts[0]
                    evidence_map[feat_id] = {
                        "test_path": parts[3],
                        "commit": parts[6]
                    }
    return evidence_map

def find_matrix_sheet(wb):
    for name in ["Feature Matrix", "Master Roadmap", "Features", "Sheet1"]:
        if name in wb.sheetnames:
            return wb[name], name
    return wb[wb.sheetnames[0]], wb.sheetnames[0]

def detect_feature_column(ws):
    header_row = 1
    feat_col = 1
    for r in range(1, 10):
        for c in range(1, 15):
            val = str(ws.cell(r, c).value or "").strip().lower()
            if val in ["feature", "feature name", "task / deliverable", "task", "work item"]:
                header_row = r
                feat_col = c
                return header_row, feat_col
    return header_row, feat_col

def process_workbook(wb_path: Path, evidence_map: dict, git_commit: str, timestamp: str):
    wb = openpyxl.load_workbook(wb_path)
    slot_num = int(wb_path.name.split("_")[0])
    
    # Load features.jsonl for this slot
    slug_matches = list((ROOT_DIR / "projects").glob(f"{slot_num:02d}-*"))
    if not slug_matches:
        return
    slug = slug_matches[0].name
    
    jsonl_path = ROOT_DIR / "projects" / slug / "spec" / "features.jsonl"
    if not jsonl_path.exists():
        return
        
    with open(jsonl_path) as f:
        features = [json.loads(l) for l in f if l.strip()]
        
    feat_by_name = {f.get("feature", f.get("name", "")).strip(): f for f in features}
    feat_by_id = {f["id"].strip(): f for f in features}

    new_headers = ["Feature ID", "Implemented", "Verified", "Status", "Evidence", "Evidence Commit"]

    # 1. Update Primary Matrix Sheet
    ws, sheet_name = find_matrix_sheet(wb)
    header_row, feat_col = detect_feature_column(ws)
    
    # Find last content column in header row
    last_col = ws.max_column
    while last_col > 1 and ws.cell(header_row, last_col).value is None:
        last_col -= 1
    start_col = last_col + 1

    for idx, h in enumerate(new_headers):
        cell = ws.cell(header_row, start_col + idx, h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")

    feat_idx = 1
    matrix_verified_col = start_col + 2
    matrix_imp_col = start_col + 1
    matrix_status_col = start_col + 3
    matrix_first_row = header_row + 1
    matrix_last_row = header_row

    for r in range(header_row + 1, ws.max_row + 1):
        feat_name_cell = ws.cell(r, feat_col).value
        if not feat_name_cell or not str(feat_name_cell).strip():
            continue
        feat_name = str(feat_name_cell).strip()
        feat_id = f"{slot_num:02d}-{feat_idx:03d}"
        feat = feat_by_id.get(feat_id, feat_by_name.get(feat_name, {}))
        
        status = feat.get("status", "Specified")
        is_verified = status == "Verified"
        is_implemented = status in ["Implemented", "Verified"]
        
        ev = evidence_map.get(feat_id, {})
        ev_test = ev.get("test_path", "")
        ev_commit = ev.get("commit", "") if ev_test else ""

        # Write cells
        c_id = ws.cell(r, start_col, feat_id)
        c_imp = ws.cell(r, start_col + 1, "✓" if is_implemented else "✗")
        c_ver = ws.cell(r, start_col + 2, "✓" if is_verified else "✗")
        c_stat = ws.cell(r, start_col + 3, status)
        c_ev = ws.cell(r, start_col + 4, ev_test)
        c_com = ws.cell(r, start_col + 5, ev_commit)

        for cell, is_tick in [(c_imp, is_implemented), (c_ver, is_verified)]:
            cell.font = GREEN_FONT if is_tick else RED_FONT
            cell.fill = GREEN_FILL if is_tick else RED_FILL
            cell.alignment = Alignment(horizontal="center")
            cell.border = BORDER_THIN
            
        c_id.border = BORDER_THIN
        c_stat.border = BORDER_THIN
        c_ev.border = BORDER_THIN
        c_com.border = BORDER_THIN
        
        matrix_last_row = r
        feat_idx += 1

    # 2. Update Detailed Analysis Sheet (if present)
    if "Detailed Analysis" in wb.sheetnames:
        ws_det = wb["Detailed Analysis"]
        det_header_row, det_feat_col = detect_feature_column(ws_det)
        det_last_col = ws_det.max_column
        while det_last_col > 1 and ws_det.cell(det_header_row, det_last_col).value is None:
            det_last_col -= 1
        det_start_col = det_last_col + 1

        for idx, h in enumerate(new_headers):
            cell = ws_det.cell(det_header_row, det_start_col + idx, h)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(horizontal="center", vertical="center")

        det_feat_idx = 1
        for r in range(det_header_row + 1, ws_det.max_row + 1):
            feat_name_cell = ws_det.cell(r, det_feat_col).value
            if not feat_name_cell or not str(feat_name_cell).strip():
                continue
            feat_name = str(feat_name_cell).strip()
            feat_id = f"{slot_num:02d}-{det_feat_idx:03d}"
            feat = feat_by_id.get(feat_id, feat_by_name.get(feat_name, {}))
            
            status = feat.get("status", "Specified")
            is_verified = status == "Verified"
            is_implemented = status in ["Implemented", "Verified"]
            ev = evidence_map.get(feat_id, {})
            ev_test = ev.get("test_path", "")
            ev_commit = ev.get("commit", "") if ev_test else ""

            c_id = ws_det.cell(r, det_start_col, feat_id)
            c_imp = ws_det.cell(r, det_start_col + 1, "✓" if is_implemented else "✗")
            c_ver = ws_det.cell(r, det_start_col + 2, "✓" if is_verified else "✗")
            c_stat = ws_det.cell(r, det_start_col + 3, status)
            c_ev = ws_det.cell(r, det_start_col + 4, ev_test)
            c_com = ws_det.cell(r, det_start_col + 5, ev_commit)

            for cell, is_tick in [(c_imp, is_implemented), (c_ver, is_verified)]:
                cell.font = GREEN_FONT if is_tick else RED_FONT
                cell.fill = GREEN_FILL if is_tick else RED_FILL
                cell.alignment = Alignment(horizontal="center")
                cell.border = BORDER_THIN
                
            c_id.border = BORDER_THIN
            c_stat.border = BORDER_THIN
            c_ev.border = BORDER_THIN
            c_com.border = BORDER_THIN
            det_feat_idx += 1

    # 3. Add 'Implementation Status' sheet with live rollup formulas
    if "Implementation Status" in wb.sheetnames:
        del wb["Implementation Status"]
    ws_status = wb.create_sheet(title="Implementation Status")

    # Header Title Banner
    ws_status.merge_cells("A1:G1")
    title_cell = ws_status["A1"]
    title_cell.value = f"IMPLEMENTATION & VERIFICATION STATUS — PROJECT {slot_num:02d} ({slug})"
    title_cell.font = Font(name="Arial", size=13, bold=True, color="FFFFFF")
    title_cell.fill = PatternFill(start_color="1F4E78", fill_type="solid")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")

    # Metadata
    ws_status["A2"] = "Generated At:"
    ws_status["B2"] = timestamp
    ws_status["A3"] = "Git Commit:"
    ws_status["B3"] = git_commit
    ws_status["A4"] = "Source File:"
    ws_status["B4"] = wb_path.name
    for r in range(2, 5):
        ws_status[f"A{r}"].font = Font(name="Arial", bold=True)

    # Rollup Table Header
    ws_status["A6"] = "METRIC"
    ws_status["B6"] = "TOTAL"
    ws_status["C6"] = "IMPLEMENTED (✓)"
    ws_status["D6"] = "VERIFIED (✓)"
    ws_status["E6"] = "% VERIFIED"
    for col in ["A", "B", "C", "D", "E"]:
        c = ws_status[f"{col}6"]
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center")

    # Formulas referencing primary sheet
    matrix_ref = f"'{sheet_name}'!"
    ver_col_letter = openpyxl.utils.get_column_letter(matrix_verified_col)
    imp_col_letter = openpyxl.utils.get_column_letter(matrix_imp_col)
    
    ws_status["A7"] = "All Features"
    ws_status["B7"] = f"=COUNTA({matrix_ref}{ver_col_letter}{matrix_first_row}:{ver_col_letter}{matrix_last_row})"
    ws_status["C7"] = f'=COUNTIF({matrix_ref}{imp_col_letter}{matrix_first_row}:{imp_col_letter}{matrix_last_row}, "✓")'
    ws_status["D7"] = f'=COUNTIF({matrix_ref}{ver_col_letter}{matrix_first_row}:{ver_col_letter}{matrix_last_row}, "✓")'
    ws_status["E7"] = "=D7/B7"
    ws_status["E7"].number_format = "0.0%"

    # P0 Burndown Header
    ws_status["A10"] = "P0 BURNDOWN LIST"
    ws_status["A10"].font = Font(name="Arial", size=11, bold=True, color="1F4E78")

    ws_status["A11"] = "Feature ID"
    ws_status["B11"] = "Feature Name"
    ws_status["C11"] = "Priority"
    ws_status["D11"] = "Implemented"
    ws_status["E11"] = "Verified"
    ws_status["F11"] = "Status"
    ws_status["G11"] = "Evidence Test / Artifact"
    for col in ["A", "B", "C", "D", "E", "F", "G"]:
        c = ws_status[f"{col}11"]
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center")

    p0_row = 12
    for feat in features:
        if feat.get("priority") == "P0":
            fid = feat["id"]
            name = feat.get("feature", feat.get("name", ""))
            status = feat.get("status", "Specified")
            is_ver = status == "Verified"
            is_imp = status in ["Implemented", "Verified"]
            ev = evidence_map.get(fid, {})
            ev_test = ev.get("test_path", "")

            ws_status[f"A{p0_row}"] = fid
            ws_status[f"B{p0_row}"] = name
            ws_status[f"C{p0_row}"] = "P0"
            
            c_imp = ws_status[f"D{p0_row}"]
            c_imp.value = "✓" if is_imp else "✗"
            c_imp.font = GREEN_FONT if is_imp else RED_FONT
            c_imp.fill = GREEN_FILL if is_imp else RED_FILL
            c_imp.alignment = Alignment(horizontal="center")

            c_ver = ws_status[f"E{p0_row}"]
            c_ver.value = "✓" if is_ver else "✗"
            c_ver.font = GREEN_FONT if is_ver else RED_FONT
            c_ver.fill = GREEN_FILL if is_ver else RED_FILL
            c_ver.alignment = Alignment(horizontal="center")

            ws_status[f"F{p0_row}"] = status
            ws_status[f"G{p0_row}"] = ev_test
            
            for col in ["A", "B", "C", "D", "E", "F", "G"]:
                ws_status[f"{col}{p0_row}"].border = BORDER_THIN
            p0_row += 1

    # Adjust column widths
    ws_status.column_dimensions["A"].width = 16
    ws_status.column_dimensions["B"].width = 50
    ws_status.column_dimensions["C"].width = 14
    ws_status.column_dimensions["D"].width = 16
    ws_status.column_dimensions["E"].width = 16
    ws_status.column_dimensions["F"].width = 16
    ws_status.column_dimensions["G"].width = 65

    # Save to evidence/workbooks/
    out_dir = ROOT_DIR / "evidence" / "workbooks"
    out_dir.mkdir(parents=True, exist_ok=True)
    out_name = f"{wb_path.stem}_STATUS.xlsx"
    out_path = out_dir / out_name
    wb.save(out_path)
    print(f"✓ Generated status workbook: {out_path}")

def main():
    print("=== Generating Excel Status Workbooks with Live Verification Ticks ===")
    evidence_map = load_evidence_map()
    
    try:
        git_commit = subprocess.check_output(["git", "rev-parse", "--short", "HEAD"], cwd=ROOT_DIR).decode().strip()
    except Exception:
        git_commit = "unknown"
        
    timestamp = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    
    spec_files = sorted([f for f in (ROOT_DIR / "inputs" / "specs").glob("*.xlsx") if f.name.startswith(("04_", "05_", "06_"))])
    for sf in spec_files:
        process_workbook(sf, evidence_map, git_commit, timestamp)

if __name__ == "__main__":
    main()
