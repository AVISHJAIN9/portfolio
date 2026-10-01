#!/usr/bin/env python3
import openpyxl
import glob
import re
import os

def main():
    feats = {}
    for slot, fname in [("04", "04_Columnar_Query_Engine_ULTIMATE_145_Master.xlsx"), ("05", "05_Mini_QUIC_ULTIMATE_145_Master.xlsx")]:
        wb = openpyxl.load_workbook(f"inputs/specs/{fname}", data_only=True)
        ws = wb["Detailed Analysis"]
        f_col = 2 if slot == "04" else 1
        p_col = 7 if slot == "04" else 6
        idx = 1
        for r in range(2, ws.max_row + 1):
            name = ws.cell(r, f_col).value
            prio = ws.cell(r, p_col).value
            if name:
                fid = f"{slot}-{idx:03d}"
                feats[fid] = {"name": str(name).strip(), "priority": str(prio).strip()}
                idx += 1

    feat_re = re.compile(r"//\s*FEAT:\s*([0-9]{2}-[0-9]{3})")
    fn_re = re.compile(r"fn\s+([a-zA-Z0-9_]+)\s*\(")

    for slot, sname in [("04", "04-columnar-query-engine"), ("05", "05-mini-quic")]:
        tdir = f"projects/{sname}/backend/tests"
        print(f"\n=======================================================")
        print(f"AUDIT OF TEST ANNOTATIONS FOR PROJECT {slot} ({sname})")
        print(f"=======================================================")
        print(f"{'FEAT ID':<8} | {'Workbook Feature Name':<45} | {'Prio':<4} | {'Test Function':<45} | Valid?")
        print("-" * 125)
        for tf in sorted(glob.glob(f"{tdir}/*.rs")):
            curr_feat = None
            with open(tf) as fp:
                for line in fp:
                    m = feat_re.search(line)
                    if m:
                        curr_feat = m.group(1)
                    m2 = fn_re.search(line)
                    if m2 and curr_feat:
                        fn_name = m2.group(1)
                        info = feats.get(curr_feat, {"name": "UNKNOWN", "priority": "?"})
                        feat_name = info["name"]
                        prio = info["priority"]
                        
                        # Check semantic match
                        is_valid = "PASS"
                        if curr_feat == "05-015" and "checksum" in fn_name.lower():
                            is_valid = "FAIL (05-015 is X25519, test is checksum)"
                        elif curr_feat == "04-044" and "query_cache" in fn_name.lower():
                            is_valid = "FAIL (04-044 is Plan cache, test is result cache)"
                        elif curr_feat == "05-002" and "short_header" in fn_name.lower():
                            is_valid = "FAIL (05-002 is Long-header, 05-003 is Short-header)"
                            
                        print(f"{curr_feat:<8} | {feat_name:<45} | {prio:<4} | {fn_name:<45} | {is_valid}")
                        curr_feat = None

if __name__ == "__main__":
    main()
